"""Audita as planilhas anuais do INSS por município (Ministério da Previdência) e mede a distorção
do órgão pagador por escala. Não ingere nem publica dados.

Uso: python -I scripts/auditar_inss_municipios.py PASTA_DOS_XLSX . SAIDA.json
Avaliação e hashes dos arquivos: research/avaliacao-inss-beneficios-municipios.md
"""
import json, re, statistics, sys, unicodedata
from collections import defaultdict
import openpyxl

RAW, REPO = sys.argv[1], sys.argv[2]
YEARS = range(2019, 2026)
COLS = ["aposent_total", "aposent_idade", "aposent_invalidez", "aposent_tempo", "pensoes", "auxilios", "outros_prev", "total_prev", "assist_ble", "total"]
ibge = json.load(open(REPO + "/data/reference/ibge_municipios_raw.json"))
REF = {f"{m['id']:07d}": m for m in ibge}
UF_SIGLA = {m["regiao-imediata"]["regiao-intermediaria"]["UF"]["nome"]: m["regiao-imediata"]["regiao-intermediaria"]["UF"]["sigla"] for m in ibge}
norm = lambda s: unicodedata.normalize("NFD", str(s)).encode("ascii", "ignore").decode().lower()

def pop_num(v):
    """População: inteiro, ou texto como '461.748(1)' (milhar com ponto + marcador de nota)."""
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, str):
        m = re.fullmatch(r"\s*([\d.]+)\s*(\(\d+\))?\s*", v)
        if m: return float(m.group(1).replace(".", ""))
    return None

def code_col(rows, h):
    """Coluna do código decidida pelos dados (o cabeçalho de 2024 está invertido)."""
    for r in rows[h + 1:h + 40]:
        for j in (0, 1):
            if r[j] is not None and re.fullmatch(r"\d{7}", str(r[j]).strip().split(".")[0]): return j
    raise ValueError("coluna de código não encontrada")

def num(v):
    if v is None or (isinstance(v, str) and not v.strip()): return None
    try: return float(v)
    except ValueError: return None  # texto de nota ou cabeçalho repetido

def parse_sheet(ws, with_pop):
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows[:12]) if any(c and "ibge" in norm(c) for c in r[:2]))
    ci = code_col(rows, h)
    ni = 1 - ci
    pop_i = next((j for j, c in enumerate(rows[h]) if c and "popula" in norm(c)), None)
    data, brasil, sem_codigo, notas = {}, None, [], []
    for r in rows[h + 3:]:
        code, name = r[ci], r[ni]
        vals = [num(v) for v in r[3:13]]
        if "brasil" in {norm(code).strip(), norm(name).strip()} and all(v is not None for v in vals):
            brasil = dict(zip(COLS, vals)); continue
        if code is not None and re.fullmatch(r"\d{7}", str(code).strip().split(".")[0]):
            c = str(code).strip().split(".")[0]
            data[c] = dict(nome=name, uf=r[2], **dict(zip(COLS, vals)), pop=pop_num(r[pop_i]) if with_pop and pop_i is not None else None)
        elif name and all(v is not None for v in vals):
            sem_codigo.append((name, r[2], vals[-1]))
        elif code is not None or name is not None:
            notas.append(str(code if code is not None else name)[:160])
    return data, brasil, sem_codigo, notas

def checks(name, data, brasil, sem_codigo, tol):
    out = {"linhas": len(data), "sem_codigo": sem_codigo}
    out["codigos_fora_ibge"] = sorted(set(data) - set(REF))
    out["ibge_ausentes"] = sorted(set(REF) - set(data))
    bad = 0
    for r in data.values():
        v = {k: r[k] for k in COLS}
        if any(x is None for x in v.values()): bad += 1; continue
        if abs(v["aposent_idade"] + v["aposent_invalidez"] + v["aposent_tempo"] - v["aposent_total"]) > tol: bad += 1
        elif abs(v["aposent_total"] + v["pensoes"] + v["auxilios"] + v["outros_prev"] - v["total_prev"]) > tol: bad += 1
        elif abs(v["total_prev"] + v["assist_ble"] - v["total"]) > tol: bad += 1
    out["erros_soma_interna"] = bad
    out["vazios"] = sum(1 for r in data.values() for k in COLS if r[k] is None)
    if brasil:
        soma = sum(r["total"] or 0 for r in data.values()) + sum(s[2] for s in sem_codigo)
        out["brasil_total"] = brasil["total"]; out["brasil_menos_soma"] = round(brasil["total"] - soma, 2)
    else:
        out["brasil_total"] = None
    return out

result, series = {}, {}
for y in YEARS:
    wb = openpyxl.load_workbook(f"{RAW}/ben_municipios_especie_{y}.xlsx", read_only=True, data_only=True)
    names = wb.sheetnames
    pick = lambda pred: next(n for n in names if pred(norm(n)))
    sq = pick(lambda n: n.startswith("qtd") or n.startswith("quantidade"))
    sv = pick(lambda n: "valor" in n and ("dez" in n) and "medio" not in n)
    sa = pick(lambda n: "valor" in n and (f"total_{y}" in n or n.endswith(f"- {y}")))
    titles = {}
    for n in names:
        rows = list(wb[n].iter_rows(values_only=True, max_row=4))
        titles[n] = next((str(r[0]) for r in rows if r[0] and len(str(r[0])) > 30), "")
    yr = {"abas": names, "titulos": {n: t[:200] for n, t in titles.items()}}
    q = parse_sheet(wb[sq], True); v = parse_sheet(wb[sv], False); a = parse_sheet(wb[sa], False)
    yr["qtd_dez"] = checks("qtd", *q[:3], 0); yr["valor_dez"] = checks("val", *v[:3], 1); yr["valor_ano"] = checks("ano", *a[:3], 1)
    # Zeros explícitos e população
    yr["zeros_total"] = sum(1 for r in q[0].values() if r["total"] == 0)
    yr["pop_ausente"] = sum(1 for r in q[0].values() if not r["pop"])
    try:
        wc = openpyxl.load_workbook(f"{RAW}/ben_municipios_clientela_{y}.xlsx", read_only=True, data_only=True)
        rows = list(wc.worksheets[0].iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows[:12]) if any(c and "ibge" in norm(c) for c in r[:2]))
        ci = code_col(rows, h)
        qi = next(j for j, c in enumerate(rows[h]) if c and norm(c).startswith("quantidade"))
        ai = [j for j, c in enumerate(rows[h]) if c and norm(c).startswith("valor")][-1]
        cl = {}
        for r in rows[h + 2:]:
            c = str(r[ci]).strip().split(".")[0] if r[ci] is not None else ""
            if re.fullmatch(r"\d{7}", c): cl[c] = (num(r[qi]), num(r[qi + 1]), num(r[qi + 2]), num(r[ai]), num(r[ai + 1]), num(r[ai + 2]))
        iguais = sum(1 for c, t in cl.items() if c in q[0] and t[0] == q[0][c]["total_prev"])
        urb_rur = sum(1 for t in cl.values() if None not in t[:3] and t[1] + t[2] != t[0])
        ano_igual = sum(1 for c, t in cl.items() if c in a[0] and t[3] is not None and abs(t[3] - a[0][c]["total_prev"]) <= 1)
        ano_urbano = sum(1 for c, t in cl.items() if c in a[0] and t[4] is not None and abs(t[4] - a[0][c]["total_prev"]) <= 1)
        yr["clientela"] = {"linhas": len(cl), "qtd_dez_igual_especie": iguais, "urbano_mais_rural_diferente": urb_rur,
                           "valor_ano_igual_total_prev": ano_igual, "valor_ano_igual_so_urbano": ano_urbano}
    except FileNotFoundError:
        yr["clientela"] = None
    result[y] = yr
    series[y] = {c: (r["total"], r["pop"], r["aposent_total"]) for c, r in q[0].items()}

def dissim(units, key_of):
    """Índice de dissimilaridade: % dos benefícios que precisaria mudar de unidade, dentro de cada grupo,
    para ficar proporcional à população. Limite superior da distorção: inclui diferenças etárias reais."""
    groups = defaultdict(list)
    for c, (b, p) in units.items(): groups[key_of(c)].append((b, p))
    moved = total = 0
    for g in groups.values():
        B = sum(b for b, _ in g); P = sum(p for _, p in g)
        if B == 0 or P == 0: continue
        moved += B * 0.5 * sum(abs(b / B - p / P) for b, p in g); total += B
    return 100 * moved / total

def dist(values):
    s = sorted(values); n = len(s)
    q = lambda p: s[min(n - 1, int(p * (n - 1)))]
    return {"n": n, "min": round(s[0], 1), "p5": round(q(.05), 1), "mediana": round(q(.5), 1), "p95": round(q(.95), 1), "max": round(s[-1], 1),
            "abaixo_5": sum(1 for x in s if x < 5), "acima_40": sum(1 for x in s if x > 40)}

ri = lambda c: REF[c]["regiao-imediata"]["id"]
uf = lambda c: REF[c]["regiao-imediata"]["regiao-intermediaria"]["UF"]["sigla"]
escalas = {}
for y, s in series.items():
    mun = {c: (b, p) for c, (b, p, _) in s.items() if c in REF and b is not None and p}
    agg = lambda key: {k: (sum(b for c, (b, _) in mun.items() if key(c) == k), sum(p for c, (_, p) in mun.items() if key(c) == k)) for k in {key(c) for c in mun}}
    reg, est = agg(ri), agg(uf)
    escalas[y] = {
        "beneficios_por_100_hab": {"municipio": dist([100 * b / p for b, p in mun.values()]),
                                     "regiao_imediata": dist([100 * b / p for b, p in reg.values()]),
                                     "uf": dist([100 * b / p for b, p in est.values()])},
        "dissimilaridade_pct": {
            "municipios_dentro_da_regiao_imediata": round(dissim(mun, ri), 1),
            "municipios_dentro_da_uf": round(dissim(mun, uf), 1),
            "regioes_imediatas_dentro_da_uf": round(dissim(reg, lambda k: str(k)[:2]), 1),
            "ufs_no_brasil": round(dissim(est, lambda k: "BR"), 1)},
    }

# Estabilidade ano a ano do polo: municípios com muito poucos benefícios continuam assim?
baixos = {y: {c for c, (b, p, _) in s.items() if p and b is not None and 100 * b / p < 2} for y, s in series.items()}
estab = {y: len(baixos[y] & baixos[y - 1]) for y in YEARS if y - 1 in baixos}
json.dump({"validacao": result, "escalas": escalas, "abaixo_2_por_100": {y: len(v) for y, v in baixos.items()}, "abaixo_2_persistentes": estab},
          open(sys.argv[3], "w"), ensure_ascii=False, indent=1, default=str)
print("ok")
